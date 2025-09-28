from fastapi import APIRouter, UploadFile, File, HTTPException
from pathlib import Path
import shutil
import pypdfium2 as pdfium
from uuid import uuid4
import pandas as pd
import os
from typing import List, Dict, Tuple, Any, Optional
from openpyxl import load_workbook
from openpyxl.styles import PatternFill
from openpyxl.chart import BarChart, Reference  # kept in case you re-enable later
from PIL import Image
import zipfile
import re
import matplotlib
matplotlib.use("Agg")  # ensure headless
import matplotlib.pyplot as plt

from functions import ocr  # <-- your YOLO+OCR wrapper
from functions.chart_utils import ChartUtils
from functions.excel_utils import ExcelUtils
from functions.grade_utils import GradeUtils
from functions.matnum_utils import MatNumUtils
from functions.ocr_post_utils import OCRPostUtils
from functions.page_utils import PageUtils
from functions.pdf_utils import PdfUtils
from functions.question_utils import QuestionUtils
from functions.student_utils import StudentUtils

router = APIRouter()
OUTPUT_DIR = Path("processed_results")
OUTPUT_DIR.mkdir(parents=True, exist_ok=True)


# --------------------------
# Configurable grading table (generic)
# --------------------------
DEFAULT_GRADING_TABLE = [
    (0,  "5,0"),
    (50, "4,0"),
    (55, "3,7"),
    (60, "3,3"),
    (65, "3,0"),
    (70, "2,7"),
    (75, "2,3"),
    (80, "2,0"),
    (85, "1,7"),
    (90, "1,3"),
    (95, "1,0"),
]

# Colors (hex without alpha for openpyxl)
COLORS = {
    "full":        "C6EFCE",  # light green
    "zero":        "C00000",  # deep red (changed from FFC7CE)
    "no_match":    "EE82EE",  # violet
    "double_match":"800080",  # purple
    "invalid":     "FFA07A",  # orange-like for invalid/plausibility
    "normalized":  "FFA500",  # orange for normalized (distinct)
    "page_error":  "D3D3D3"   # light gray for page plausibility error
}

# Font color overrides for readability
FONT_COLORS = {
    "double_match": "FFFFFF",  # white on purple
    "zero": "FFFFFF"  # white on deep red
}


@router.post("/process-image/")
async def process_image_from_upload(file: UploadFile = File(...)):
    try:
        unique_id = uuid4().hex
        pdf_folder = OUTPUT_DIR / unique_id
        pdf_folder.mkdir(parents=True, exist_ok=True)

        saved_pdf_path = pdf_folder / file.filename
        with open(saved_pdf_path, "wb") as buffer:
            shutil.copyfileobj(file.file, buffer)

        # Convert PDF to images and run YOLO+OCR per page
        pdf = pdfium.PdfDocument(str(saved_pdf_path))
        pages_results: List[Dict[str, Any]] = []

        for i in range(len(pdf)):
            page_num = i + 1
            page_folder = pdf_folder / f"image_{page_num}"
            page_folder.mkdir(parents=True, exist_ok=True)

            page = pdf[i]
            pil_image = page.render(scale=4).to_pil()
            page_image_path = page_folder / "page.jpg"
            pil_image.save(page_image_path)

            # call your YOLO+OCR wrapper - should return dict with "results" list
            result = ocr.process_and_ocr_image(str(page_image_path), output_dir=page_folder)
            result.update({"page": page_num, "page_folder": str(page_folder)})
            pages_results.append(result)

        # Split pages into students by Mat_num
        student_groups = MatNumUtils.split_pages_by_matnum(pages_results)

        students_rows: List[Dict[str, Any]] = []
        students_norm_flags: List[Dict[str, bool]] = []
        students_numeric: List[Dict[str, Optional[float]]] = []
        students_qmaps: List[Dict[str, Any]] = []
        students_seen: List[Dict[str, int]] = []
        students_issues: List[Dict[str, Any]] = []
        per_student_folders: List[Path] = []

        # process each student group
        for idx, pages in enumerate(student_groups):
            student_folder = pdf_folder / f"student_{idx+1}"
            student_folder.mkdir(parents=True, exist_ok=True)

            student_info, qmap, per_q_seen, page_markers = StudentUtils.extract_from_pages(pages)
            page_check_msg, page_ok = PageUtils.page_plausibility_check(page_markers, pages)

            # build student row
            row, norm_flags, numeric_achieved = StudentUtils.build_student_row_and_flags(student_info, qmap, pdf_folder, page_check_msg)
            # compute and attach totals/percent/mark already done in build_student_row_and_flags

            issues, per_q_status = QuestionUtils.run_plausibility_checks(qmap, numeric_achieved, per_q_seen)
            # append page-related plausibility issue to issues
            if not page_ok:
                issues.append(f"Page check: {page_check_msg}")

            # collect
            students_rows.append(row)
            students_norm_flags.append(norm_flags)
            students_numeric.append(numeric_achieved)
            students_qmaps.append(qmap)
            students_seen.append(per_q_seen)
            students_issues.append({"issues": issues, "per_q_status": per_q_status, "page_check": page_check_msg})

            # create per-student annotated PDF (detected.jpg pages for this group)
            imgs = []
            for p in pages:
                pf_path = Path(p.get("page_folder"))
                det = pf_path / "detected.jpg"
                if det.exists():
                    imgs.append(Image.open(det).convert("RGB"))
            if imgs:
                student_pdf_path = student_folder / "annotated_student.pdf"
                imgs[0].save(student_pdf_path, save_all=True, append_images=imgs[1:])
            else:
                # fallback: create PDF from page.jpg
                fallback_imgs = []
                for p in pages:
                    pf_path = Path(p.get("page_folder"))
                    pj = pf_path / "page.jpg"
                    if pj.exists():
                        fallback_imgs.append(Image.open(pj).convert("RGB"))
                if fallback_imgs:
                    fallback_path = student_folder / "annotated_student.pdf"
                    fallback_imgs[0].save(fallback_path, save_all=True, append_images=fallback_imgs[1:])

            # create per-student excel (with primus sheet included) - NO charts inside
            per_excel = student_folder / "student_result.xlsx"
            StudentUtils.save_students_excel_and_primus([row], [norm_flags], [numeric_achieved], [qmap], [per_q_seen], per_excel,unique_id)

            per_student_folders.append(student_folder)

        # also produce a merged annotated PDF (all pages) and a combined Excel for all students
        merged_pdf = PdfUtils.save_annotated_pdf(pdf_folder, out_name="annotated_all.pdf")
        combined_excel = pdf_folder / "all_students_results.xlsx"
        StudentUtils.save_students_excel_and_primus(students_rows, students_norm_flags, students_numeric, students_qmaps, students_seen, combined_excel,unique_id)

        # if multiple students -> bundle into one ZIP and include ONLY the overall grade chart inside the ZIP
        bundle_path: Optional[str] = None
        if len(per_student_folders) > 1:
            # create grade distribution chart
            chart_path = pdf_folder / "grade_distribution.png"
            chart_file = ChartUtils.generate_grade_distribution_chart(students_rows, chart_path)

            zip_path = pdf_folder / "batch_results.zip"
            with zipfile.ZipFile(zip_path, "w", compression=zipfile.ZIP_DEFLATED) as zf:
                # add per-student folders
                for sf in per_student_folders:
                    for root, _, files in os.walk(sf):
                        for f in files:
                            fp = Path(root) / f
                            arc = fp.relative_to(pdf_folder)
                            zf.write(fp, arc)
                # add the overall chart at root of zip (only if created)
                if chart_file and Path(chart_file).exists():
                    zf.write(chart_file, Path(chart_file).name)
            bundle_path = str(zip_path)

        response = {
            "message": "✅ Processing complete",
            "output_dir": str(pdf_folder),
            "combined_excel": str(combined_excel),
            "annotated_all_pdf": merged_pdf,
            "zip_if_batch": bundle_path,
            "students_count": len(students_rows),
            "students_issues": students_issues
        }
        return response

    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Failed to process PDF: {str(e)}")












